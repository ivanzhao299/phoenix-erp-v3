#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown → Word(.docx) 转换器（支持标题/表格/列表/加粗/分隔线，中文字体）。
用法: python3 md_to_docx.py <input.md> <output.docx> ["文档标题"]
"""
import sys, re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CJK_FONT = "微软雅黑"
HEADER_FILL = "1F4E79"   # 深蓝表头
HEADER_TEXT = RGBColor(0xFF, 0xFF, 0xFF)

def set_cjk(run):
    run.font.name = CJK_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'):
        rfonts.set(qn(a), CJK_FONT)

def shade_cell(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hexcolor)
    tcpr.append(shd)

def add_runs(paragraph, text, base_size=10.5, bold=False, color=None):
    """解析 **加粗** 与 ~~删除线~~，逐段加 run。"""
    text = text.replace('`','')
    # 按 ** 和 ~~ 切分
    tokens = re.split(r'(\*\*.*?\*\*|~~.*?~~)', text)
    for tok in tokens:
        if not tok:
            continue
        b, strike, content = bold, False, tok
        if tok.startswith('**') and tok.endswith('**'):
            b = True; content = tok[2:-2]
        elif tok.startswith('~~') and tok.endswith('~~'):
            strike = True; content = tok[2:-2]
        run = paragraph.add_run(content)
        run.font.size = Pt(base_size); run.bold = b
        if strike: run.font.strike = True
        if color: run.font.color.rgb = color
        set_cjk(run)

def parse_table(lines):
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip('|').split('|')]
        rows.append(cells)
    # 去掉分隔行（第2行 ---）
    if len(rows) >= 2 and all(re.match(r'^:?-+:?$', c.replace(' ','')) for c in rows[1] if c):
        header = rows[0]; body = rows[2:]
    else:
        header = rows[0]; body = rows[1:]
    return header, body

def add_table(doc, header, body):
    ncol = len(header)
    table = doc.add_table(rows=1, cols=ncol)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    # 表头
    hcells = table.rows[0].cells
    for i, h in enumerate(header):
        shade_cell(hcells[i], HEADER_FILL)
        p = hcells[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, h, base_size=10, bold=True, color=HEADER_TEXT)
    # 表体
    for r in body:
        cells = table.add_row().cells
        for i in range(ncol):
            val = r[i] if i < len(r) else ''
            p = cells[i].paragraphs[0]
            add_runs(p, val, base_size=9.5)
    doc.add_paragraph()

def main():
    src, out = sys.argv[1], sys.argv[2]
    title = sys.argv[3] if len(sys.argv) > 3 else None
    with open(src, encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()
    # 默认样式字体
    normal = doc.styles['Normal']
    normal.font.name = CJK_FONT; normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), CJK_FONT)

    if title:
        h = doc.add_heading('', level=0)
        add_runs(h, title, base_size=20, bold=True, color=RGBColor(0x1F,0x4E,0x79))
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1; continue
        # 表格
        if s.startswith('|') and i+1 < len(lines) and lines[i+1].strip().startswith('|'):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                tbl.append(lines[i]); i += 1
            header, body = parse_table(tbl)
            add_table(doc, header, body)
            continue
        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            level = len(m.group(1)); txt = m.group(2)
            hp = doc.add_heading('', level=min(level,4))
            sizes = {1:18,2:15,3:13,4:11.5,5:11,6:10.5}
            add_runs(hp, txt, base_size=sizes.get(level,12), bold=True,
                     color=RGBColor(0x1F,0x4E,0x79))
            i += 1; continue
        # 分隔线
        if re.match(r'^(-{3,}|\*{3,}|_{3,})$', s):
            p = doc.add_paragraph(); pr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'6')
            bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),'BBBBBB')
            pbdr.append(bottom); pr.append(pbdr)
            i += 1; continue
        # 引用
        if s.startswith('>'):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(18)
            add_runs(p, s.lstrip('> ').strip(), base_size=9.5, color=RGBColor(0x66,0x66,0x66))
            i += 1; continue
        # 无序列表
        m = re.match(r'^[-*]\s+(.*)$', s)
        if m:
            p = doc.add_paragraph(style='List Bullet'); add_runs(p, m.group(1)); i += 1; continue
        # 有序列表
        m = re.match(r'^\d+\.\s+(.*)$', s)
        if m:
            p = doc.add_paragraph(style='List Number'); add_runs(p, m.group(1)); i += 1; continue
        # 普通段落
        p = doc.add_paragraph(); add_runs(p, s)
        i += 1

    doc.save(out)
    print(f"已生成: {out}")

if __name__ == '__main__':
    main()
